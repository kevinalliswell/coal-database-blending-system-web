from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from backend.utils.excel_utils import create_unified_a4_report
from test_result_snapshots import client, solve


def workbook_text(workbook):
    return '\n'.join(str(cell.value) for sheet in workbook for row in sheet.iter_rows()
                     for cell in row if cell.value is not None)


def test_report_exposes_completed_run_basis_heat_price_and_assay_provenance(client):
    payload = solve(client)
    response = client.post('/api/optimization/export/current', json=payload)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data))
    text = workbook_text(workbook)
    assert '恒容低位' in text and 'Qnet,v' in text
    assert '元/收到吨' in text and '收到基' in text
    assert 'CBOS-assay-2.0' in text and 'testadmin' in text
    assert '化验原值与采用值' in workbook.sheetnames
    assert '换算与运行快照' in workbook.sheetnames
    sheet = workbook['化验原值与采用值']
    headers = [cell.value for cell in sheet[1]]
    rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    first_id = payload['input_snapshot']['assays'][0]['id']
    ash = next(row for row in rows if row['化验修订ID'] == first_id and row['指标'] == 'ash_content')
    assert ash['原始值'] == 10
    assert ash['原始基准'] == 'd'
    assert ash['采用值'] == pytest.approx(9.2)
    assert ash['采用基准'] == 'ar'
    fixed_carbon = next(row for row in rows if row['化验修订ID'] == first_id and row['指标'] == 'fixed_carbon')
    assert fixed_carbon['原始值'] is None
    assert fixed_carbon['采用值'] is not None
    assert '派生' in fixed_carbon['来源说明']


def test_report_uses_snapshot_type_even_if_current_or_top_level_metadata_disagrees(client):
    payload = solve(client)
    payload['input_snapshot']['heat_type'] = 'gr_v'
    payload['input_snapshot']['ratio_basis'] = 'd'
    payload['result']['data_validation']['ratio_basis'] = 'ar'
    response = client.post('/api/optimization/export/current', json=payload)
    text = workbook_text(load_workbook(BytesIO(response.data)))
    assert '恒容高位' in text and 'Qgr,v' in text
    assert '元/干吨' in text


def test_legacy_missing_metadata_is_explicitly_unrecorded_not_current_defaults():
    workbook = Workbook()
    create_unified_a4_report(workbook, {'result': {'coal_names': [], 'optimal_ratios': []}})
    text = workbook_text(workbook)
    assert '热值定义：未记录' in text
    assert '计价口径：未记录' in text
    assert '公式版本：未记录' in text
    assert '原始化验及采用值快照未记录' in text
    assert '恒容低位' not in text


def test_spreadsheet_provenance_labels_are_text_not_executable_formulas(client):
    payload = solve(client)
    payload['input_snapshot']['assays'][0]['label'] = '=1+1'
    response = client.post('/api/optimization/export/current', json=payload)
    workbook = load_workbook(BytesIO(response.data), data_only=False)
    for name in ('化验原值与采用值', '换算与运行快照'):
        assert all(cell.data_type != 'f' for row in workbook[name] for cell in row)
