#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from datetime import datetime
from io import BytesIO
import json
import logging

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


TITLE_FILL = PatternFill('solid', fgColor='DCE6F1')
SECTION_FILL = PatternFill('solid', fgColor='EAF2F8')
TABLE_HEADER_FILL = PatternFill('solid', fgColor='5B9BD5')
LABEL_FILL = PatternFill('solid', fgColor='F7F9FB')
SUCCESS_FILL = PatternFill('solid', fgColor='E2F0D9')
WARNING_FILL = PatternFill('solid', fgColor='FFF2CC')
ERROR_FILL = PatternFill('solid', fgColor='FDE9E7')

THIN_SIDE = Side(style='thin', color='B7C0CC')
MEDIUM_SIDE = Side(style='medium', color='7F8C9A')
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
MEDIUM_BORDER = Border(left=MEDIUM_SIDE, right=MEDIUM_SIDE, top=MEDIUM_SIDE, bottom=MEDIUM_SIDE)

TITLE_FONT = Font(name='Microsoft YaHei', size=18, bold=True, color='1F1F1F')
SUBTITLE_FONT = Font(name='Microsoft YaHei', size=10, color='5A6573')
SECTION_FONT = Font(name='Microsoft YaHei', size=12, bold=True, color='1F4E79')
TABLE_HEADER_FONT = Font(name='Microsoft YaHei', size=10, bold=True, color='FFFFFF')
LABEL_FONT = Font(name='Microsoft YaHei', size=10, bold=True, color='334155')
VALUE_FONT = Font(name='Microsoft YaHei', size=10, color='111827')
SMALL_FONT = Font(name='Microsoft YaHei', size=9, color='475569')
STATUS_OK_FONT = Font(name='Microsoft YaHei', size=10, bold=True, color='2E7D32')
STATUS_WARN_FONT = Font(name='Microsoft YaHei', size=10, bold=True, color='B7791F')
STATUS_BAD_FONT = Font(name='Microsoft YaHei', size=10, bold=True, color='C0392B')


def create_unified_a4_report(wb, data, user_constraints=None):
    """创建适合 A4 打印的统一优化报告。"""
    ws = wb.active
    ws.title = '喷煤混配优化报告'
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A5'

    result = _extract_result(data)
    constraints = user_constraints or data.get('constraints', {}) or {}

    _setup_a4_page(ws)
    _setup_columns(ws)

    row = 1
    row = _write_title_block(ws, row, result)
    snapshot = data.get('input_snapshot') or {}
    row = _write_data_validation_block(ws, row, result, snapshot)
    row = _write_summary_block(ws, row, result)
    row = _write_ratio_block(ws, row, result)
    row = _write_performance_block(ws, row, result, constraints)
    row = _write_constraints_block(ws, row, result, constraints)
    row = _write_conclusion_block(ws, row, result)

    last_row = max(row - 1, 1)
    ws.print_area = f'A1:H{last_row}'
    ws.print_title_rows = '$1:$4'

    ws.oddHeader.center.text = '喷煤混配优化报告'
    ws.oddFooter.left.text = '&D &T'
    ws.oddFooter.center.text = '第 &P 页 / 共 &N 页'
    ws.oddFooter.right.text = '智能煤炭混配系统'
    _write_provenance_sheets(wb, snapshot)


def _extract_result(data):
    if isinstance(data, dict) and isinstance(data.get('result'), dict):
        return data['result']
    return data or {}


def _setup_a4_page(ws):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.45
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2


def _setup_columns(ws):
    widths = {
        'A': 13,
        'B': 15,
        'C': 13,
        'D': 14,
        'E': 13,
        'F': 15,
        'G': 13,
        'H': 15,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def _write_title_block(ws, row, result):
    _merge_and_write(
        ws,
        f'A{row}:H{row}',
        '喷煤工艺混配优化结果报告',
        TITLE_FONT,
        TITLE_FILL,
        MEDIUM_BORDER,
        Alignment(horizontal='center', vertical='center'),
    )
    ws.row_dimensions[row].height = 28
    row += 1

    objective_label = result.get('optimization_target') or '多目标Pareto最优'

    subtitle = (
        f"生成时间：{_display_timestamp(result.get('timestamp'))}    "
        f"|    优化目标：{objective_label}    "
        f"|    数学可行性：{'满足数值约束' if result.get('is_feasible') else '未形成可行方案'}"
    )
    _merge_and_write(
        ws,
        f'A{row}:H{row}',
        subtitle,
        SUBTITLE_FONT,
        None,
        THIN_BORDER,
        Alignment(horizontal='center', vertical='center'),
    )
    ws.row_dimensions[row].height = 20
    row += 2
    return row


def _data_validation_label(result):
    status = (result.get('data_validation') or {}).get('status')
    if status == 'unverified':
        return '基准未确认，仅用于筛选'
    if status == 'verified':
        return '基准检查通过（不代表现场工艺验证）'
    return '基准验证未记录'


def _snapshot_metadata(snapshot, validation):
    basis = snapshot.get('ratio_basis') or validation.get('ratio_basis')
    basis_label = {'ar': '收到基 ar（收到煤质量）', 'd': '干基 d（干煤质量）',
                   'ad': '空干基 ad（分析样质量）'}.get(basis, '未记录')
    price_label = {'ar': '元/收到吨', 'd': '元/干吨', 'ad': '元/空干吨'}.get(basis, '未记录')
    heat = {'gr_v': '恒容高位 Qgr,v', 'net_v': '恒容低位 Qnet,v',
            'net_p': '恒压低位 Qnet,p'}.get(snapshot.get('heat_type'), '未记录')
    operator = snapshot.get('operator') or snapshot.get('actor') or {}
    operator_label = ' / '.join(str(operator[key]) for key in ('display_name', 'username', 'role') if operator.get(key)) or '未记录'
    return [f'配比质量基准：{basis_label}', f'热值定义：{heat}',
            f'计价口径：{price_label}',
            f'公式版本：{snapshot.get("formula_version") or "未记录"}',
            f'设置版本：{(snapshot.get("settings") or {}).get("version", "未记录")}',
            f'操作者：{operator_label}',
            ('热值单位：kcal_IT/kg；1 kcal_IT = 4.1868 kJ'
             if snapshot.get('schema_version') == 2 else '热值数值单位：kcal/kg；旧记录卡定义未记录')]


def _write_data_validation_block(ws, row, result, snapshot=None):
    snapshot = snapshot or {}
    validation = result.get('data_validation') or {}
    row = _write_section_title(ws, row, '数据基准与用途')
    lines = [
        _data_validation_label(result),
        *_snapshot_metadata(snapshot, validation),
        ('原始化验及采用值见“化验原值与采用值”工作表，换算和运行参数见“换算与运行快照”。'
         if snapshot.get('assays') else '原始化验及采用值快照未记录；不以当前资料或默认设置补写旧历史。'),
        '检查范围为输入基准一致性及工业分析闭合，不代表现场工艺验证。',
        *(validation.get('warnings') or []),
    ]
    for line in lines:
        _merge_and_write(ws, f'A{row}:H{row}', line, VALUE_FONT, WARNING_FILL,
                         THIN_BORDER, Alignment(horizontal='left', vertical='center', wrap_text=True))
        ws.row_dimensions[row].height = max(24, 16 * ((len(line) // 50) + 1))
        row += 1
    return row + 1


def _append_provenance_row(ws, values):
    # User-supplied labels/report identifiers must remain text in spreadsheets.
    safe = ["'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value
            for value in values]
    ws.append(safe)


def _style_provenance_sheet(ws):
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    _setup_a4_page(ws)
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.print_title_rows = '$1:$1'
    for cell in ws[1]:
        cell.font, cell.fill = TABLE_HEADER_FONT, TABLE_HEADER_FILL
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = SMALL_FONT
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for column in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(column)].width = 22 if ws.max_column > 4 else (32 if column < 3 else 100)
    ws.print_area = ws.dimensions


def _write_provenance_sheets(wb, snapshot):
    assays = snapshot.get('assays') or []
    if not assays:
        return
    values_sheet = wb.create_sheet('化验原值与采用值')
    _append_provenance_row(values_sheet, ['化验修订ID', '煤种/批次/修订', '指标', '原始值', '原始单位', '原始基准',
                                       '原始热值/氢定义', '采用值', '采用单位', '采用基准', '来源说明'])
    trace_sheet = wb.create_sheet('换算与运行快照')
    _append_provenance_row(trace_sheet, ['类别', '标识', '保存内容'])
    for line in _snapshot_metadata(snapshot, snapshot.get('data_validation') or {}):
        _append_provenance_row(trace_sheet, ['本次口径', '', line])
    for key in ('constraints', 'ga_config', 'algorithm_params', 'settings'):
        for field, value in (snapshot.get(key) or {}).items():
            _append_provenance_row(trace_sheet, [key, field, json.dumps(value, ensure_ascii=False)])
    for assay in assays:
        raw = assay.get('raw') or {}
        eligibility = (assay.get('eligibility') or {}).get('blending') or {}
        normalized = eligibility.get('values') or {}
        label = assay.get('label') or str(assay.get('id', '未记录'))
        for field in dict.fromkeys([*raw, *normalized]):
            item = raw.get(field) or {}
            definition = item.get('heat_type', '')
            if field == 'hydrogen_content':
                includes = item.get('hydrogen_includes_moisture')
                definition = {True: '包含水中氢', False: '不含水中氢'}.get(includes, '水中氢定义未记录')
            unit = ('kcal_IT/kg' if field == 'calorific_value' else 'CNY/t' if field == 'price' else '%') if field in normalized else None
            source = '原始值换算/直接采用' if item.get('value') is not None else '固定碳标准差减派生（原始报告未提供）' if field == 'fixed_carbon' and field in normalized else '原始值未记录'
            _append_provenance_row(values_sheet, [assay.get('id'), label, field, item.get('value'), item.get('unit'), item.get('basis'),
                definition, normalized.get(field), unit, snapshot.get('ratio_basis') if field in normalized else None, source])
        for key in ('coal_id', 'batch_id', 'revision', 'report_number', 'sampled_at', 'source', 'created_by', 'created_at'):
            _append_provenance_row(trace_sheet, ['化验身份', f'{label} / {key}', assay.get(key)])
        for line in eligibility.get('trace') or []:
            _append_provenance_row(trace_sheet, ['换算过程', label, str(line)])
    _style_provenance_sheet(values_sheet)
    _style_provenance_sheet(trace_sheet)


def _write_summary_block(ws, row, result):
    row = _write_section_title(ws, row, '一、结果概览')

    left_items = [
        ('吨煤成本', _format_currency(result.get('total_price'))),
        ('综合热值', _format_number(result.get('total_heat'), 0, 'kcal/kg')),
        ('综合灰分', _format_number(result.get('total_ash'), 2, '%')),
        ('综合挥发分', _format_number(result.get('total_volatile'), 2, '%')),
    ]
    right_items = [
        ('综合固定碳', _format_number(result.get('total_fixed_carbon'), 2, '%')),
        ('综合硫分', _format_number(result.get('total_sulfur'), 3, '%')),
        ('迭代次数', _format_int(result.get('iterations'))),
        ('计算耗时', _format_number(result.get('time_taken'), 2, '秒')),
    ]

    max_len = max(len(left_items), len(right_items))
    for index in range(max_len):
        left_label, left_value = left_items[index]
        right_label, right_value = right_items[index]
        _write_label_value_row(ws, row + index, 'A', left_label, left_value, 'E', right_label, right_value)

    status_row = row + max_len
    _write_label_value_row(
        ws,
        status_row,
        'A',
        '当前结论',
        _data_validation_label(result) if result.get('constraints_satisfied') else '需调整约束或煤种后再优化',
        'E',
        '约束满足',
        '是' if result.get('constraints_satisfied') else '否',
    )
    ws[f'B{status_row}'].font = STATUS_WARN_FONT
    ws[f'F{status_row}'].font = STATUS_OK_FONT if result.get('constraints_satisfied') else STATUS_BAD_FONT
    return status_row + 2


def _write_ratio_block(ws, row, result):
    row = _write_section_title(ws, row, '二、推荐配比')

    headers = ['序号', '煤种名称', '配比(%)', '配比(小数)', '配比角色', '备注']
    _write_table_header(ws, row, headers)
    row += 1

    ratios = list(zip(result.get('coal_names', []), result.get('optimal_ratios', [])))
    ratios = [(name, ratio) for name, ratio in ratios if ratio and ratio > 0]
    ratios.sort(key=lambda item: item[1], reverse=True)

    if not ratios:
        _merge_and_write(
            ws,
            f'A{row}:H{row}',
            '当前没有可用于打印的有效配比结果。',
            SMALL_FONT,
            LABEL_FILL,
            THIN_BORDER,
            Alignment(horizontal='center', vertical='center'),
        )
        ws.row_dimensions[row].height = 22
        return row + 2

    for index, (name, ratio) in enumerate(ratios, start=1):
        ratio_percent = ratio * 100
        values = [
            index,
            name,
            f'{ratio_percent:.2f}%',
            f'{ratio:.4f}',
            get_ratio_grade(ratio_percent),
            get_ratio_comment(ratio_percent),
        ]
        _write_table_row(ws, row, values, highlight_col=2)
        row += 1

    return row + 1


def _write_performance_block(ws, row, result, constraints):
    row = _write_section_title(ws, row, '三、关键指标与目标区间')

    headers = ['指标', '结果值', '单位', '目标区间', '状态', '评价']
    _write_table_header(ws, row, headers)
    row += 1

    performance_rows = [
        ('吨煤成本', result.get('total_price'), '元/吨', _range_text(constraints.get('price_min'), constraints.get('price_max')), _within_range(result.get('total_price'), constraints.get('price_min'), constraints.get('price_max')), evaluate_cost(result.get('total_price', 0))),
        ('综合热值', result.get('total_heat'), 'kcal/kg', _range_text(constraints.get('heat_min'), constraints.get('heat_max')), _within_range(result.get('total_heat'), constraints.get('heat_min'), constraints.get('heat_max')), evaluate_heat_value(result.get('total_heat', 0))),
        ('综合灰分', result.get('total_ash'), '%', _range_text(constraints.get('ash_min'), constraints.get('ash_max')), _within_range(result.get('total_ash'), constraints.get('ash_min'), constraints.get('ash_max')), evaluate_ash_content(result.get('total_ash', 0))),
        ('综合挥发分', result.get('total_volatile'), '%', _range_text(constraints.get('volatile_min'), constraints.get('volatile_max')), _within_range(result.get('total_volatile'), constraints.get('volatile_min'), constraints.get('volatile_max')), evaluate_volatile_matter(result.get('total_volatile', 0))),
        ('综合固定碳', result.get('total_fixed_carbon'), '%', _range_text(constraints.get('fixed_carbon_min'), constraints.get('fixed_carbon_max')), _within_range(result.get('total_fixed_carbon'), constraints.get('fixed_carbon_min'), constraints.get('fixed_carbon_max')), evaluate_fixed_carbon(result.get('total_fixed_carbon', 0))),
        ('综合硫分', result.get('total_sulfur'), '%', _range_text(constraints.get('sulfur_min'), constraints.get('sulfur_max')), _within_range(result.get('total_sulfur'), constraints.get('sulfur_min'), constraints.get('sulfur_max')), evaluate_sulfur_content(result.get('total_sulfur', 0))),
    ]

    for indicator, actual, unit, target, ok, review in performance_rows:
        status_text = '未校验' if ok is None else '达标' if ok else '偏离'
        if ok and (result.get('data_validation') or {}).get('status') == 'unverified':
            status_text = '试算数值满足'
        values = [
            indicator,
            _format_number(actual, 2 if unit == '%' or unit == '元/吨' else 0, ''),
            unit,
            target,
            status_text,
            review,
        ]
        _write_table_row(ws, row, values)
        ws[f'E{row}'].font = SMALL_FONT if ok is None else STATUS_OK_FONT if ok else STATUS_BAD_FONT
        ws[f'E{row}'].fill = LABEL_FILL if ok is None else SUCCESS_FILL if ok else ERROR_FILL
        row += 1

    return row + 1


def _write_constraints_block(ws, row, result, constraints):
    row = _write_section_title(ws, row, '四、约束检查与诊断')

    _merge_and_write(
        ws,
        f'A{row}:D{row}',
        '活跃约束 / 满足情况',
        TABLE_HEADER_FONT,
        TABLE_HEADER_FILL,
        THIN_BORDER,
        Alignment(horizontal='center', vertical='center'),
    )
    _merge_and_write(
        ws,
        f'E{row}:H{row}',
        '违反约束 / 诊断信息',
        TABLE_HEADER_FONT,
        TABLE_HEADER_FILL,
        THIN_BORDER,
        Alignment(horizontal='center', vertical='center'),
    )
    ws.row_dimensions[row].height = 22
    row += 1

    active_constraints = result.get('active_constraints', []) or []
    violation_items = result.get('constraint_violations', []) or []
    diagnosis_items = result.get('diagnosis', []) or []

    left_lines = (
        ['全部数值约束满足'] if result.get('constraints_satisfied')
        else ['当前方案未满足全部数值约束']
    )
    if active_constraints:
        left_lines.extend([f'活跃约束：{item}' for item in active_constraints[:6]])

    right_lines = []
    if violation_items:
        right_lines.extend([_constraint_text(item) for item in violation_items[:6]])
    elif diagnosis_items:
        right_lines.extend(diagnosis_items[:6])
    else:
        right_lines.append('暂无诊断信息')

    max_rows = max(len(left_lines), len(right_lines), 1)
    for offset in range(max_rows):
        left_text = left_lines[offset] if offset < len(left_lines) else ''
        right_text = right_lines[offset] if offset < len(right_lines) else ''
        _merge_and_write(
            ws,
            f'A{row}:D{row}',
            left_text,
            VALUE_FONT if offset else (STATUS_OK_FONT if result.get('constraints_satisfied') else STATUS_BAD_FONT),
            SUCCESS_FILL if result.get('constraints_satisfied') else WARNING_FILL if offset == 0 else None,
            THIN_BORDER,
            Alignment(horizontal='left', vertical='center', wrap_text=True),
        )
        _merge_and_write(
            ws,
            f'E{row}:H{row}',
            right_text,
            VALUE_FONT,
            ERROR_FILL if violation_items else LABEL_FILL,
            THIN_BORDER,
            Alignment(horizontal='left', vertical='center', wrap_text=True),
        )
        ws.row_dimensions[row].height = 22 if len(left_text) < 30 and len(right_text) < 30 else 34
        row += 1

    row += 1
    row = _write_section_title(ws, row, '五、配煤约束与算法信息')
    _write_label_value_row(
        ws,
        row,
        'A',
        '目标函数',
        result.get('optimization_target') or '多目标Pareto最优',
        'E',
        '质量偏离分',
        _format_number(result.get('fitness_value'), 4, ''),
    )
    _write_label_value_row(
        ws,
        row + 1,
        'A',
        '约束违约度',
        _format_number(result.get('violation_score'), 4, ''),
        'E',
        '算法模型',
        '遗传算法 + 线性可行性预检查',
    )
    return row + 3


def _write_conclusion_block(ws, row, result):
    row = _write_section_title(ws, row, '六、结论与打印说明')

    conclusion = _build_conclusion_text(result)
    _merge_and_write(
        ws,
        f'A{row}:H{row + 3}',
        conclusion,
        VALUE_FONT,
        LABEL_FILL,
        MEDIUM_BORDER,
        Alignment(horizontal='left', vertical='top', wrap_text=True),
    )
    for current in range(row, row + 4):
        ws.row_dimensions[current].height = 24
    return row + 5


def _write_section_title(ws, row, title):
    _merge_and_write(
        ws,
        f'A{row}:H{row}',
        title,
        SECTION_FONT,
        SECTION_FILL,
        THIN_BORDER,
        Alignment(horizontal='left', vertical='center'),
    )
    ws.row_dimensions[row].height = 22
    return row + 1


def _write_label_value_row(ws, row, label_col_1, label_1, value_1, label_col_2, label_2, value_2):
    value_col_1 = get_column_letter(ord(label_col_1) - 64 + 1)
    value_col_2 = get_column_letter(ord(label_col_2) - 64 + 1)

    _write_cell(ws, f'{label_col_1}{row}', label_1, LABEL_FONT, LABEL_FILL, THIN_BORDER, Alignment(horizontal='left', vertical='center'))
    _write_cell(ws, f'{value_col_1}{row}', value_1, VALUE_FONT, None, THIN_BORDER, Alignment(horizontal='left', vertical='center'))
    _write_cell(ws, f'{label_col_2}{row}', label_2, LABEL_FONT, LABEL_FILL, THIN_BORDER, Alignment(horizontal='left', vertical='center'))
    _write_cell(ws, f'{value_col_2}{row}', value_2, VALUE_FONT, None, THIN_BORDER, Alignment(horizontal='left', vertical='center'))

    ws.row_dimensions[row].height = 20


def _write_table_header(ws, row, headers):
    columns = ['A', 'B', 'C', 'D', 'E', 'F']
    for column, value in zip(columns, headers):
        _write_cell(
            ws,
            f'{column}{row}',
            value,
            TABLE_HEADER_FONT,
            TABLE_HEADER_FILL,
            THIN_BORDER,
            Alignment(horizontal='center', vertical='center'),
        )
    ws.row_dimensions[row].height = 22


def _write_table_row(ws, row, values, highlight_col=None):
    columns = ['A', 'B', 'C', 'D', 'E', 'F']
    for index, (column, value) in enumerate(zip(columns, values), start=1):
        fill = LABEL_FILL if highlight_col and index == highlight_col else None
        alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        if column in {'B', 'F'}:
            alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        _write_cell(ws, f'{column}{row}', value, VALUE_FONT, fill, THIN_BORDER, alignment)
    ws.row_dimensions[row].height = 20


def _merge_and_write(ws, merged_range, value, font, fill, border, alignment):
    ws.merge_cells(merged_range)
    cells = ws[merged_range]
    for row in cells:
        for cell in row:
            cell.fill = fill or PatternFill(fill_type=None)
            cell.border = border
            cell.alignment = alignment

    cell = ws[merged_range.split(':')[0]]
    cell.value = value
    cell.font = font
    cell.fill = fill or PatternFill(fill_type=None)
    cell.border = border
    cell.alignment = alignment


def _write_cell(ws, reference, value, font, fill, border, alignment):
    cell = ws[reference]
    cell.value = value
    cell.font = font
    cell.fill = fill or PatternFill(fill_type=None)
    cell.border = border
    cell.alignment = alignment


def _display_timestamp(timestamp):
    if not timestamp:
        return datetime.now().strftime('%Y-%m-%d %H:%M')
    try:
        return datetime.fromisoformat(str(timestamp)).strftime('%Y-%m-%d %H:%M')
    except (TypeError, ValueError):
        return str(timestamp)


def _range_text(min_value, max_value):
    if min_value is None and max_value is None:
        return '--'
    if min_value is None:
        return f'<= {max_value}'
    if max_value is None:
        return f'>= {min_value}'
    return f'{min_value} ~ {max_value}'


def _within_range(value, min_value=None, max_value=None):
    if min_value is None and max_value is None:
        return None
    if value is None:
        return False
    if min_value is not None and value < min_value:
        return False
    if max_value is not None and value > max_value:
        return False
    return True


def _format_currency(value):
    if value is None:
        return '--'
    return f'¥{float(value):.2f}/吨'


def _format_number(value, decimals=2, suffix=''):
    if value is None:
        return '--'
    return f'{float(value):.{decimals}f}{suffix}'


def _format_int(value):
    if value is None:
        return '--'
    return str(int(value))


def _constraint_text(item):
    if isinstance(item, str):
        return item
    if isinstance(item, dict):
        return item.get('message') or item.get('detail') or item.get('label') or str(item)
    return str(item)


def _build_conclusion_text(result):
    feasible = bool(result.get('constraints_satisfied'))
    total_price = _format_currency(result.get('total_price'))
    total_heat = _format_number(result.get('total_heat'), 0, ' kcal/kg')
    active_constraints = result.get('active_constraints', []) or []

    if (result.get('data_validation') or {}).get('status') == 'unverified':
        return (
            '本次为未确认数据的试算结果：基准未确认，仅用于筛选，不可直接用于生产。'
            f'数值约束{"满足" if feasible else "未满足"}；吨煤成本 {total_price}，综合热值 {total_heat}。'
            '请先确认检验数据、热值种类及质量/价格/配比基准，再开展工艺验证。'
        )

    if feasible:
        constraint_text = '；'.join(active_constraints[:3]) if active_constraints else '当前方案存在一定约束余量'
        return (
            f'本次优化形成数值约束下的候选方案，不代表现场工艺验证。'
            f'在满足硬约束前提下，当前方案的吨煤成本为 {total_price}，综合热值为 {total_heat}。'
            f'需要重点关注的活跃约束为：{constraint_text}。'
            f'导出时已按 A4 纵向页面布局优化，适合直接打印留档或用于会议汇报。'
        )

    diagnosis = result.get('diagnosis', []) or ['当前未形成可行方案，请结合约束设置与候选煤种重新评估。']
    return (
        f'本次优化尚未形成满足全部硬约束的最终方案。'
        f'当前结果可作为约束调参与煤种筛选的参考，不建议直接用于生产执行。'
        f'主要诊断信息：{"；".join(diagnosis[:3])}。'
        f'报表已按 A4 打印版式整理，便于线下会签、复盘和继续优化。'
    )


def get_ratio_grade(ratio_percent):
    """根据配比比例评级。"""
    if ratio_percent >= 40:
        return '主配煤'
    if ratio_percent >= 20:
        return '重要组分'
    if ratio_percent >= 10:
        return '辅助组分'
    return '微量补充'


def get_ratio_comment(ratio_percent):
    """根据配比比例输出说明。"""
    if ratio_percent >= 50:
        return '主导煤种'
    if ratio_percent >= 30:
        return '主要成本/质量贡献来源'
    if ratio_percent >= 15:
        return '承担结构性调节作用'
    if ratio_percent >= 5:
        return '用于精细校准'
    return '占比很小，可重点复核必要性'


def create_coal_template():
    """创建煤种数据导入模板。"""
    template_data = pd.DataFrame({
        '序号': [1],
        '工业分析基准': ['unknown'],
        '热值基准': ['unknown'],
        '热值类型': ['unknown'],
        '计价基准': ['unknown'],
        '煤种名称': ['示例煤种'],
        '煤种类型': ['烟煤'],
        '供应商/产地': ['示例供应商'],
        '价格(元/吨)': [500.0],
        '热值(kcal/kg)': [6000],
        '灰分(%)': [15.0],
        '挥发分(%)': [25.0],
        '固定碳(%)': [60.0],
        '全硫(%)': [0.6],
        '全水分(%)': [10.0],
        '可磨性HGI': [60],
        '燃尽率(%)': [95.0],
        '含碳量C(%)': [70.0],
        '含氢量H(%)': [4.0],
        '含氮量N(%)': [1.0],
    })

    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        template_data.to_excel(writer, sheet_name='煤种数据模板', index=False)
    output.seek(0)
    return output


def create_coal_template_with_mad():
    """创建包含空干基水分字段的煤种数据导入模板。"""
    template_data = pd.DataFrame({
        '序号': [1],
        '工业分析基准': ['unknown'],
        '热值基准': ['unknown'],
        '热值类型': ['unknown'],
        '计价基准': ['unknown'],
        '煤种名称': ['示例煤种'],
        '煤种类型': ['烟煤'],
        '供应商/产地': ['示例供应商'],
        '价格(元/吨)': [500.0],
        '热值(kcal/kg)': [6000],
        '灰分(%)': [15.0],
        '挥发分(%)': [25.0],
        '固定碳(%)': [60.0],
        '全硫(%)': [0.6],
        '全水分(%)': [10.0],
        '空干基水分(%)': [8.0],
        '可磨性HGI': [60],
        '燃尽率(%)': [95.0],
        '含碳量C(%)': [70.0],
        '含氢量H(%)': [4.0],
        '含氮量N(%)': [1.0],
    })

    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        template_data.to_excel(writer, sheet_name='煤种数据模板', index=False)
    output.seek(0)
    return output


def evaluate_cost(cost):
    """评价成本水平。"""
    if cost <= 500:
        return '优秀'
    if cost <= 600:
        return '良好'
    if cost <= 700:
        return '中等'
    return '偏高'


def evaluate_heat_value(heat):
    """评价热值水平。"""
    if heat >= 6300:
        return '优秀'
    if heat >= 6000:
        return '良好'
    if heat >= 5800:
        return '中等'
    return '偏低'


def evaluate_ash_content(ash):
    """评价灰分水平。"""
    if ash <= 10:
        return '优秀'
    if ash <= 15:
        return '良好'
    if ash <= 18:
        return '中等'
    return '偏高'


def evaluate_volatile_matter(volatile):
    """评价挥发分水平。"""
    if 28 <= volatile <= 32:
        return '优秀'
    if 25 <= volatile <= 35:
        return '良好'
    if 15 <= volatile <= 40:
        return '中等'
    return '需复核'


def evaluate_fixed_carbon(carbon):
    """评价固定碳水平。"""
    if 65 <= carbon <= 70:
        return '优秀'
    if 60 <= carbon <= 75:
        return '良好'
    if 55 <= carbon <= 80:
        return '中等'
    return '需复核'


def evaluate_sulfur_content(sulfur):
    """评价硫分水平。"""
    if sulfur <= 0.5:
        return '优秀'
    if sulfur <= 0.6:
        return '良好'
    if sulfur <= 0.8:
        return '中等'
    return '偏高'
