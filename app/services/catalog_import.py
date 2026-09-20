"""Explicit column mapping and single-use, server-held import snapshots."""
from copy import deepcopy
import csv
from datetime import date, datetime, timedelta, timezone
import hashlib
import io
import math
import secrets

from openpyxl import Workbook, load_workbook
from sqlalchemy import func

from app.extensions import db
from app.models.coal import Coal
from app.models.catalog import AssayRevision, CoalBatch, ImportPreview
from app.services.catalog_service import (METRICS, CatalogConflict, create_assay, create_batch,
                                          create_coal, date_field, text_field)

MAX_ROWS = 2000
MAX_BYTES = 10 * 1024 * 1024
IDENTITY_FIELDS = ('coal_name', 'type', 'batch_code', 'supplier', 'received_at', 'report_number', 'sampled_at', 'notes')
INFO_FIELDS = ('coal_id', 'batch_id', 'assay_revision_id', 'revision', 'source', 'archived', 'created_at')
METRIC_FIELDS = tuple(f'{metric}.{part}' for metric in METRICS for part in ('value', 'unit', 'basis')) + (
    'calorific_value.heat_type', 'hydrogen_content.hydrogen_includes_moisture')
FIELDS = IDENTITY_FIELDS + METRIC_FIELDS
HEADERS = IDENTITY_FIELDS + INFO_FIELDS + METRIC_FIELDS
ALIASES = {'煤种名称': 'coal_name', 'name': 'coal_name', '煤种类型': 'type',
           '批次编号': 'batch_code', '批次': 'batch_code', '供应商': 'supplier',
           '收货日期': 'received_at', '报告编号': 'report_number', '采样日期': 'sampled_at', '备注': 'notes',
           '灰分': 'ash_content.value', '灰分(%)': 'ash_content.value',
           '挥发分': 'volatile_matter.value', '挥发分(%)': 'volatile_matter.value',
           '固定碳': 'fixed_carbon.value', '固定碳(%)': 'fixed_carbon.value',
           '全硫': 'total_sulfur.value', '全硫(%)': 'total_sulfur.value',
           '热值': 'calorific_value.value', '热值(kcal/kg)': 'calorific_value.value',
           '全水分': 'total_moisture.value', '空干基水分': 'air_dry_moisture.value',
           '价格': 'price.value', '价格(元/吨)': 'price.value', 'HGI': 'hgi.value',
           '含碳量': 'carbon_content.value', '含氢量': 'hydrogen_content.value',
           '含氮量': 'nitrogen_content.value', 'SiO2': 'sio2_content.value'}
ALIASES.update({metric: f'{metric}.value' for metric in METRICS})


def _now():
    # SQLAlchemy's portable DateTime column stores naive UTC on SQLite.
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _table(text=None, file_bytes=None, filename=None):
    if file_bytes is not None:
        if not isinstance(file_bytes, bytes) or len(file_bytes) > MAX_BYTES:
            raise ValueError('文件不能超过10MB')
        suffix = (filename or '').lower().rsplit('.', 1)[-1]
        if suffix == 'xlsx':
            try:
                book = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=False)
                try:
                    if len(book.worksheets) != 1:
                        raise ValueError('请仅保留一个数据工作表后导入')
                    sheet = book.worksheets[0]
                    if sheet.max_column > 128:
                        raise ValueError('最多支持128列')
                    rows = []
                    for row in sheet.iter_rows(values_only=True):
                        rows.append(list(row))
                        if len(rows) > MAX_ROWS + 1:
                            raise ValueError('每次最多导入2000行')
                finally:
                    book.close()
            except ValueError:
                raise
            except Exception:
                raise ValueError('无法读取XLSX文件') from None
        elif suffix == 'csv':
            for encoding in ('utf-8-sig', 'gb18030'):
                try:
                    text = file_bytes.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise ValueError('CSV编码须为UTF-8或GB18030')
            rows = None
        else:
            raise ValueError('仅支持CSV或XLSX文件')
    else:
        rows = None
    if rows is None:
        if not isinstance(text, str) or len(text.encode('utf-8')) > MAX_BYTES:
            raise ValueError('请提供不超过10MB的CSV或制表符分隔文本')
        first_line = text.splitlines()[0] if text.splitlines() else ''
        delimiter = '\t' if '\t' in first_line else ','
        rows = []
        try:
            for row in csv.reader(io.StringIO(text.lstrip('\ufeff')), delimiter=delimiter, strict=True):
                if len(row) > 128:
                    raise ValueError('最多支持128列')
                rows.append(row)
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError('每次最多导入2000行')
        except csv.Error:
            raise ValueError('CSV/粘贴内容格式错误') from None
    if not rows:
        raise ValueError('文件内容为空')
    headers = [str(item).strip() if item is not None else '' for item in rows[0]]
    if not all(headers) or len(set(headers)) != len(headers):
        raise ValueError('表头不能为空或重复')
    body = []
    for number, values in enumerate(rows[1:], 2):
        if all(_blank(value) for value in values):
            continue
        if len(values) > len(headers):
            raise ValueError(f'第{number}行列数超过表头')
        body.append((number, dict(zip(headers, values + [None] * (len(headers) - len(values))))))
    if not body:
        raise ValueError('没有可导入的数据行')
    return headers, body


def _issue(field, code, message):
    return {'field': field, 'code': code, 'message': message}


def _number(value, field):
    if isinstance(value, bool):
        raise ValueError(f'{field}不能为布尔值')
    try:
        result = float(value)
    except (ValueError, TypeError):
        raise ValueError(f'{field}必须为数字') from None
    if not math.isfinite(result):
        raise ValueError(f'{field}必须为有限数字')
    return result


def _boolean(value, field):
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        if value.strip().lower() in ('true', '是', '1'):
            return True
        if value.strip().lower() in ('false', '否', '0'):
            return False
        if value.strip().lower() in ('unknown', 'null', ''):
            return None
    raise ValueError(f'{field}必须为true、false或空白')


def _identity(row):
    return (row['coal_name'], row['batch_code'])


def _target_state(row):
    coal = Coal.query.filter_by(name=row['coal_name']).first()
    if coal is None:
        return {'coal_id': None, 'batch_id': None, 'revision': 0}
    if coal.type != row['type']:
        raise ValueError('现有同名煤种类型不同，请更正名称或类型')
    if coal.archived:
        raise ValueError('煤种已归档')
    batch = CoalBatch.query.filter_by(coal_id=coal.id, code=row['batch_code']).first()
    if batch is None:
        return {'coal_id': coal.id, 'batch_id': None, 'revision': 0}
    if batch.archived:
        raise ValueError('批次已归档')
    for field in ('supplier', 'received_at'):
        if row.get(field) and row[field] != getattr(batch, field):
            raise ValueError(f'已有批次的{field}与导入不一致；不能通过化验导入覆盖批次')
    revision = db.session.query(func.max(AssayRevision.revision)).filter_by(batch_id=batch.id).scalar() or 0
    return {'coal_id': coal.id, 'batch_id': batch.id, 'revision': revision}


def preview_import(*, user_id, settings, text=None, file_bytes=None, filename=None,
                   mapping=None, metric_metadata=None, confirmed_headers=None):
    from app.services.assay_conversion import validate_raw_assay, evaluate_assay
    headers, body = _table(text, file_bytes, filename)
    mapping = {} if mapping is None else mapping
    metadata = {} if metric_metadata is None else metric_metadata
    confirmed = [] if confirmed_headers is None else confirmed_headers
    if not isinstance(mapping, dict) or set(mapping) - set(headers):
        raise ValueError('mapping须为当前表头到目标字段的对象')
    if not isinstance(metadata, dict) or set(metadata) - set(METRICS):
        raise ValueError('metric_metadata包含未知指标')
    if not isinstance(confirmed, list) or any(not isinstance(value, str) for value in confirmed) or set(confirmed) - set(headers):
        raise ValueError('confirmed_headers须为当前表头名称列表')
    for metric, values in metadata.items():
        if not isinstance(values, dict) or any(f'{metric}.{key}' not in METRIC_FIELDS or key == 'value' for key in values):
            raise ValueError('metric_metadata仅用于填写单位、基准和热值/氢定义')
    actual, info, issues = {}, [], []
    for header in headers:
        suggested = header if header in FIELDS else ALIASES.get(header)
        ambiguous = suggested is None and header not in INFO_FIELDS or (header in ALIASES and '.' in (suggested or ''))
        chosen = mapping[header] if header in mapping else suggested
        if chosen is not None and chosen not in FIELDS:
            raise ValueError(f'不支持的映射字段: {chosen}')
        if ambiguous and header not in mapping and header not in confirmed:
            issues.append(_issue(header, 'mapping_confirmation', f'请确认“{header}”的字段映射；不需要的列请明确忽略'))
        info.append({'name': header, 'suggested': suggested, 'ambiguous': bool(ambiguous)})
        actual[header] = chosen
    targets = [value for value in actual.values() if value]
    if len(set(targets)) != len(targets):
        issues.append(_issue('mapping', 'duplicate_mapping', '多个列映射到同一字段，请保留一个'))
    for field in ('coal_name', 'type', 'batch_code'):
        if field not in targets:
            issues.append(_issue(field, 'missing_column', f'缺少必需列{field}'))
    rows, coal_types, batch_metadata = [], {}, {}
    for row_number, original in body:
        row = {'row': row_number, 'raw': {}, 'issues': []}
        mapped = {actual[key]: value for key, value in original.items() if actual[key]}
        try:
            for field in IDENTITY_FIELDS:
                value = mapped.get(field)
                if isinstance(value, (datetime, date)):
                    value = value.isoformat()
                if field in ('sampled_at', 'received_at'):
                    row[field] = date_field(value, field)
                else:
                    row[field] = text_field(None if _blank(value) else str(value), field,
                                           10000 if field == 'notes' else 120 if field == 'report_number' else 50 if field == 'type' else 100,
                                           field in ('coal_name', 'type', 'batch_code'))
            for metric in METRICS:
                value = mapped.get(f'{metric}.value')
                if _blank(value):
                    continue
                details = deepcopy(metadata.get(metric, {}))
                details['value'] = _number(value, metric)
                for part in ('unit', 'basis', 'heat_type', 'hydrogen_includes_moisture'):
                    if f'{metric}.{part}' in METRIC_FIELDS and not _blank(mapped.get(f'{metric}.{part}')):
                        item = mapped[f'{metric}.{part}']
                        details[part] = _boolean(item, part) if part == 'hydrogen_includes_moisture' else str(item).strip()
                missing = [part for part in ('unit', 'basis') if not details.get(part)]
                if metric == 'calorific_value' and not details.get('heat_type'):
                    missing.append('heat_type')
                source_header = next(header for header, target in actual.items() if target == f'{metric}.value')
                if missing and source_header not in confirmed:
                    row['issues'].append(_issue(source_header, 'metadata_confirmation', f'{metric}缺少{", ".join(missing)}；请补充或确认保留未知'))
                row['raw'][metric] = details
            row['raw'] = validate_raw_assay(row['raw'])
            row['eligibility'] = evaluate_assay(row['raw'], settings)
            row['target'] = _target_state(row)
            if row['coal_name'] in coal_types and coal_types[row['coal_name']] != row['type']:
                raise ValueError('同一文件中同名煤种的类型不一致')
            coal_types[row['coal_name']] = row['type']
            identity = _identity(row)
            values = (row['supplier'], row['received_at'])
            if identity in batch_metadata and batch_metadata[identity] != values:
                raise ValueError('同一批次的供应商/收货日期在文件内不一致')
            batch_metadata[identity] = values
        except ValueError as exc:
            row['issues'].append(_issue('row', 'invalid_row', str(exc)))
        rows.append(row)
    can_commit = not issues and not any(row['issues'] for row in rows)
    token, expiry = None, None
    if can_commit:
        token = secrets.token_urlsafe(32)
        expiry = _now() + timedelta(minutes=30)
        ImportPreview.query.filter(ImportPreview.expires_at < _now()).delete(synchronize_session=False)
        db.session.add(ImportPreview(token_hash=hashlib.sha256(token.encode()).hexdigest(), user_id=user_id,
                                     payload={'rows': rows}, expires_at=expiry))
        db.session.commit()
    return {'headers': info, 'mapping': actual, 'fields': list(FIELDS), 'rows': rows, 'issues': issues,
            'can_commit': can_commit, 'preview_token': token,
            'expires_at': expiry.isoformat() + 'Z' if expiry else None}


def commit_import(token, *, user_id):
    if not isinstance(token, str) or len(token) > 200:
        raise ValueError('导入预览令牌无效')
    token_hash = hashlib.sha256(token.encode()).hexdigest()
    preview = db.session.get(ImportPreview, token_hash)
    if preview is None or preview.user_id != user_id:
        raise ValueError('导入预览不存在或不属于当前用户')
    if preview.consumed_at is not None or preview.expires_at <= _now():
        raise ValueError('导入预览已使用或过期，请重新预览')
    rows = deepcopy(preview.payload['rows'])
    # Claim within the same transaction before target checks. On SQLite the
    # write lock serializes another import; any later failure rolls it back.
    claimed = ImportPreview.query.filter_by(token_hash=token_hash, user_id=user_id, consumed_at=None).filter(
        ImportPreview.expires_at > _now()).update({'consumed_at': _now()}, synchronize_session=False)
    if claimed != 1:
        raise CatalogConflict('导入预览已被使用，请重新预览')
    checked = set()
    for row in rows:
        identity = _identity(row)
        if identity in checked:
            continue
        try:
            state = _target_state(row)
        except ValueError:
            raise CatalogConflict('预览后的煤种/批次已变更，请重新预览') from None
        if state != row['target']:
            raise CatalogConflict('预览后的煤种/批次或化验修订已变更，请重新预览')
        checked.add(identity)
    result = {'created_coals': 0, 'created_batches': 0, 'created_revisions': 0, 'assay_revision_ids': []}
    for row in rows:
        coal = Coal.query.filter_by(name=row['coal_name']).first()
        if coal is None:
            coal = create_coal({'name': row['coal_name'], 'type': row['type'], 'supplier': row['supplier']})
            result['created_coals'] += 1
        batch = CoalBatch.query.filter_by(coal_id=coal.id, code=row['batch_code']).first()
        if batch is None:
            batch = create_batch(coal.id, {'code': row['batch_code'], 'supplier': row['supplier'], 'received_at': row['received_at']})
            result['created_batches'] += 1
        record = create_assay(batch.id, {key: row[key] for key in ('raw', 'report_number', 'sampled_at', 'notes')}, user_id=user_id, source='import')
        result['created_revisions'] += 1
        result['assay_revision_ids'].append(record.id)
    return result


def export_catalog(format='xlsx', *, template=False, include_archived=True):
    if format not in ('csv', 'xlsx'):
        raise ValueError('format仅支持csv或xlsx')
    records = []
    if not template:
        for record in AssayRevision.query.order_by(AssayRevision.id).all():
            batch, coal = record.batch, record.batch.coal
            if not include_archived and (coal.archived or batch.archived):
                continue
            row = {'coal_name': coal.name, 'type': coal.type, 'batch_code': batch.code,
                   'supplier': batch.supplier, 'received_at': batch.received_at,
                   'report_number': record.report_number, 'sampled_at': record.sampled_at, 'notes': record.notes,
                   'coal_id': coal.id, 'batch_id': batch.id, 'assay_revision_id': record.id, 'revision': record.revision,
                   'source': record.source, 'archived': bool(coal.archived or batch.archived),
                   'created_at': record.created_at.isoformat() if record.created_at else None}
            for metric, details in record.raw.items():
                if details:
                    row.update({f'{metric}.{part}': value for part, value in details.items()})
            records.append([row.get(header) for header in HEADERS])
    basename = 'cbos-assay-template' if template else 'cbos-assays'
    if format == 'xlsx':
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '化验原始记录'
        for values in [HEADERS, *records]:
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Export user text as text, never as an Excel formula.
        sheet.freeze_panes = 'D2'
        sheet.auto_filter.ref = sheet.dimensions
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', basename + '.xlsx'
    output = io.StringIO(newline='')
    writer = csv.writer(output)
    writer.writerow(HEADERS)
    writer.writerows(records)
    return output.getvalue().encode('utf-8-sig'), 'text/csv; charset=utf-8', basename + '.csv'
