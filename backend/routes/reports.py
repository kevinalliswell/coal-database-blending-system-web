#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import pandas as pd
from flask import Blueprint, request, jsonify, send_file
from openpyxl import Workbook
from openpyxl.styles import Font
from io import BytesIO
from datetime import datetime
import logging

from app.extensions import db
from app.models import OptimizationHistory
from app.utils.excel_utils import create_unified_a4_report
from app.utils.report_utils import generate_blending_report_filename

logger = logging.getLogger(__name__)

# 创建报告导出蓝图
reports_bp = Blueprint('reports', __name__)

def _create_and_send_report(data: dict):
    """
    核心函数：创建Excel报告并生成Flask响应。
    """
    try:
        # 创建Excel工作簿
        wb = Workbook()
        
        # 获取用户设置的约束条件（仅在导出当前结果时存在）
        user_constraints = data.get('constraints', {})
        
        # 创建统一的A4排版报告
        create_unified_a4_report(wb, data, user_constraints)
        
        # 保存到内存中的字节流
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # 使用工具函数生成文件名
        filename = generate_blending_report_filename(data)
        
        # 返回Flask send_file响应
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )
    except Exception as e:
        logger.error(f"创建并发送报告时出错: {str(e)}")
        # 以JSON格式返回内部服务器错误
        return jsonify({'status': 'error', 'message': f'报告生成失败: {str(e)}'}), 500

@reports_bp.route('/optimization/export/<int:history_id>', methods=['GET'])
def export_optimization_report(history_id):
    """导出历史Excel报告"""
    try:
        history = db.session.get(OptimizationHistory, history_id)

        if history is None:
            return jsonify({'status': 'error', 'message': '优化记录不存在'}), 404

        data = history.to_detail_dict()['result_payload']

        # 调用核心函数处理报告生成和发送
        return _create_and_send_report(data)

    except Exception as e:
        logger.error(f"导出历史Excel报告失败: {str(e)}")
        return jsonify({'status': 'error', 'message': f'导出报告失败: {str(e)}'}), 500

@reports_bp.route('/optimization/export/current', methods=['POST'])
def export_current_optimization_report():
    """导出当前优化结果"""
    try:
        data = request.json
        if not data or 'result' not in data:
            return jsonify({'status': 'error', 'message': '没有优化结果数据'}), 400
        
        # 调用核心函数处理报告生成和发送
        return _create_and_send_report(data)

    except Exception as e:
        logger.error(f"导出当前优化结果失败: {str(e)}")
        return jsonify({'status': 'error', 'message': f'导出报告失败: {str(e)}'}), 500

@reports_bp.route('/replacement_ratio/export', methods=['POST'])
def export_replacement_ratio_result():
    """导出置换比计算结果API"""
    try:
        data = request.json
        if not data:
            return jsonify({'status': 'error', 'message': '请求数据为空'}), 400
        
        # 创建Excel工作簿
        wb = Workbook()
        ws = wb.active
        ws.title = "置换比计算结果"
        
        # 设置标题
        ws['A1'] = "煤焦置换比计算结果报告"
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:D1')
        
        # 添加基本信息
        row = 3
        ws[f'A{row}'] = "计算类型:"
        ws[f'B{row}'] = data.get('calculation_type', '')
        row += 1
        
        ws[f'A{row}'] = "计算结果:"
        ws[f'B{row}'] = data.get('result', '')
        row += 1
        
        ws[f'A{row}'] = "计算时间:"
        ws[f'B{row}'] = data.get('timestamp', '')
        row += 2
        
        # 添加详细计算过程
        ws[f'A{row}'] = "详细计算过程"
        ws[f'A{row}'].font = Font(bold=True)
        row += 1
        
        ws[f'A{row}'] = "参数名称"
        ws[f'B{row}'] = "数值"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'].font = Font(bold=True)
        row += 1
        
        # 添加计算详情
        details = data.get('details', {})
        for key, value in details.items():
            ws[f'A{row}'] = key
            ws[f'B{row}'] = value
            row += 1
        
        # 保存到内存
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # 生成文件名
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"置换比计算结果_{timestamp}.xlsx"
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )
        
    except Exception as e:
        logger.error(f"导出置换比计算结果失败: {str(e)}")
        return jsonify({'status': 'error', 'message': f'导出失败: {str(e)}'}), 500 
