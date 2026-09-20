"""Result provenance and printed constraints must match the submitted run."""
from io import BytesIO
import json

import pytest
from openpyxl import load_workbook

from app import create_app
from app.config import Config
from app.extensions import db
from app.models import Coal, OptimizationHistory


from test_v2_optimization import v2_client as client, add_assay, request_for, raw_fixture


def solve(client, calculation_mode=None, ratio_basis=None):
    first = add_assay(client, '无烟煤A', 0)
    second = add_assay(client, '烟煤B', 3)
    request = request_for([first, second])
    request['constraints'].update(heat_min=6000, heat_max=7200, heat_target=6800, sulfur_max=.7)
    request['ga_params'].update(population_size=12, max_generations=2, parallel_computing=False, parallel_workers=9)
    response = client.post('/api/optimization', json=request)
    assert response.status_code == 200, response.get_json()
    assert response.get_json()['status'] == 'success'
    return response.get_json()


def performance_row(response, label):
    assert response.status_code == 200
    ws = load_workbook(BytesIO(response.data)).active
    return next([cell.value for cell in row] for row in ws.iter_rows()
                if row[0].value == label and row[3].value is not None)


def test_response_and_history_preserve_coal_properties_and_effective_configuration(client):
    payload = solve(client)
    snapshot = payload['input_snapshot']
    assert snapshot['schema_version'] == 2
    assert snapshot['ratio_basis'] == 'ar'
    assert snapshot['data_validation'] == payload['result']['data_validation']
    assert snapshot['assays'][0]['raw']['price']['value'] == 950
    assert snapshot['assays'][0]['raw']['total_moisture']['value'] == 8
    assert snapshot['assays'][0]['raw']['ash_content']['basis'] == 'd'
    assert snapshot['algorithm_coals'][0]['price'] == pytest.approx(950)
    assert snapshot['ga_config']['parallel_workers'] == 1
    assert snapshot['ga_config']['population_size'] == 12
    assert snapshot['constraints']['ash_max'] == 20
    assert payload['constraints'] == snapshot['constraints']
    batch_id = snapshot['assays'][0]['batch_id']
    changed = raw_fixture(); changed['price']['value'] = 1234
    assert client.post(f'/api/batches/{batch_id}/assays', json={'raw': changed}).status_code == 201
    history = client.get('/api/optimization/history').get_json()['data'][0]
    detail = client.get(f'/api/optimization/history/{history["id"]}').get_json()['data']
    assert detail['result_payload']['input_snapshot'] == snapshot
    assert detail['constraints'] == snapshot['constraints']


def test_removed_screening_cannot_create_new_history(client):
    response = client.post('/api/optimization', json={'selected_coal_ids': [1,2], 'calculation_mode': 'screening'})
    assert response.status_code == 400
    assert client.get('/api/optimization/history').get_json()['data'] == []


def test_current_export_uses_constraints_from_the_completed_response(client):
    payload = solve(client)
    row = performance_row(client.post('/api/optimization/export/current', json=payload), '综合热值')
    assert [float(bound) for bound in row[3].split(' ~ ')] == [6000, 7200]
    assert row[4] == '达标'


def test_report_prefers_completed_snapshot_over_later_edited_constraints(client):
    payload = solve(client)
    payload['constraints'] = {'heat_min': 6000, 'heat_max': 6100}
    row = performance_row(client.post('/api/optimization/export/current', json=payload), '综合热值')
    assert [float(bound) for bound in row[3].split(' ~ ')] == [6000, 7200]
    assert row[4] == '达标'


def test_report_marks_a_missing_window_unchecked(client):
    payload = solve(client)
    payload.pop('constraints', None)
    payload.pop('input_snapshot', None)
    row = performance_row(client.post('/api/optimization/export/current', json=payload), '综合热值')
    assert row[3] == '--'
    assert row[4] == '未校验'


def test_historical_export_keeps_legacy_constraints_without_new_snapshot(client):
    payload = solve(client)
    payload.pop('constraints', None)
    payload.pop('input_snapshot', None)
    with client.application.app_context():
        history = OptimizationHistory.query.first()
        history.result_payload = json.dumps(payload)
        history.constraints = json.dumps({'heat_min': 6100, 'heat_max': 7100})
        db.session.commit()
        history_id = history.id
    row = performance_row(client.get(f'/api/optimization/export/{history_id}'), '综合热值')
    assert row[3] == '6100 ~ 7100'
