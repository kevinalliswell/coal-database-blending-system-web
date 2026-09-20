"""New installs remain empty; optional test tables use the real import contract."""
import io
import math
from pathlib import Path

import pytest

from test_catalog_api import client  # noqa: F401
from app import create_app
from app.extensions import db
from app.models.coal import Coal
from app.models.catalog import AssayRevision
from app.services.database_service import initialize_database


def test_web_initialization_is_empty_and_preserves_existing_records(tmp_path):
    app = create_app('production', config_overrides={
        'TESTING': True, 'SQLALCHEMY_DATABASE_URI': f'sqlite:///{tmp_path / "fresh.db"}',
        'DATA_DIR': str(tmp_path), 'UPLOAD_FOLDER': str(tmp_path / 'uploads'),
    })
    initialize_database(app)
    with app.app_context():
        assert Coal.query.count() == 0
        assert AssayRevision.query.count() == 0
        db.session.add(Coal(name='已有用户煤', type='烟煤', archived=True))
        db.session.commit()
    initialize_database(app)
    with app.app_context():
        assert [(coal.name, coal.archived) for coal in Coal.query.all()] == [('已有用户煤', True)]
        assert AssayRevision.query.count() == 0
        db.session.remove(); db.engine.dispose()


def test_legacy_auto_initialize_flag_never_loads_coal_file(tmp_path):
    from app.services.coal_data_service import ensure_coal_data_initialized
    app = create_app(config_overrides={
        'TESTING': True, 'SQLALCHEMY_DATABASE_URI': f'sqlite:///{tmp_path / "legacy.db"}',
        'DATA_DIR': str(tmp_path), 'UPLOAD_FOLDER': str(tmp_path / 'uploads'),
        'AUTO_INITIALIZE_COAL_DATA': True,
        'COAL_DATA_PATH': str(Path(__file__).parents[1] / 'data/test_coal_datas.csv'),
    })
    with app.app_context():
        ensure_coal_data_initialized()
        assert Coal.query.count() == 0
        db.session.remove(); db.engine.dispose()


@pytest.mark.parametrize('format', ['xlsx', 'csv'])
def test_optional_six_coal_file_imports_and_calculates_without_auto_writes(client, format):
    assert client.get('/api/catalog').get_json()['data'] == []
    download = client.get('/api/catalog/test-data?format=' + format)
    assert download.status_code == 200
    assert client.get('/api/catalog').get_json()['data'] == []
    preview = client.post('/api/catalog/import/preview', data={
        'file': (io.BytesIO(download.data), '测试煤.' + format)})
    assert preview.status_code == 200, preview.get_json()
    data = preview.get_json()['data']
    assert data['can_commit'], data
    assert len(data['rows']) == 6
    assert all(row['eligibility']['blending']['eligible'] and row['eligibility']['theory']['eligible'] for row in data['rows'])
    assert client.get('/api/catalog').get_json()['data'] == []
    committed = client.post('/api/catalog/import/commit', json={'preview_token': data['preview_token']})
    assert committed.status_code == 201, committed.get_json()
    assays = client.get('/api/assays').get_json()['data']
    assert len(assays) == 6
    assert sum(assay['coal_type'] == '无烟煤' for assay in assays) == 3
    assert sum(assay['coal_type'] == '烟煤' for assay in assays) == 3
    assert all('测试' in assay['coal_name'] and '非生产' in assay['notes'] for assay in assays)
    for assay in assays:
        assert assay['raw']['total_moisture']['basis'] == 'ar'
        assert assay['raw']['air_dry_moisture']['basis'] == 'ad'
        assert assay['raw']['ash_content']['basis'] == 'd'
        assert assay['raw']['volatile_matter']['basis'] == 'daf'
        assert assay['raw']['total_sulfur']['basis'] == 'd'
        assert assay['raw']['calorific_value']['heat_type'] == 'net_v'
    settings = client.get('/api/settings').get_json()['data']
    calculated = client.post('/api/optimization', json={
        'settings_version': settings['version'], 'constraints': settings['constraints'],
        'components': [{'assay_revision_id': a['id'], 'min_ratio': 5, 'max_ratio': 60} for a in assays],
        'ga_params': {'population_size': 24, 'max_generations': 8, 'tournament_size': 3, 'random_seed': 42},
    })
    assert calculated.status_code == 200, calculated.get_json()
    result = calculated.get_json()['result']['recommended_solution']
    assert result['is_feasible']
    ratios = result['optimal_ratios']
    assert len(ratios) == 6 and math.isclose(sum(ratios), 1, abs_tol=1e-8)
    assert all(.05 - 1e-8 <= ratio <= .6 + 1e-8 for ratio in ratios)
    expected_price = sum(a['raw']['price']['value'] * ratio for a, ratio in zip(assays, ratios))
    expected_heat = sum(a['raw']['calorific_value']['value'] / .0041868 * ratio for a, ratio in zip(assays, ratios))
    assert result['total_price'] == pytest.approx(expected_price)
    assert result['total_heat'] == pytest.approx(expected_heat)


def test_test_table_is_distinct_from_blank_template(client):
    from openpyxl import load_workbook
    response = client.get('/api/catalog/template?format=xlsx')
    book = load_workbook(io.BytesIO(response.data))
    assert book.active.max_row == 1
    assert client.get('/api/catalog/test-data?format=txt').status_code == 400
    assert client.get('/api/catalog').get_json()['data'] == []
